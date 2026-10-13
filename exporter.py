"""
Exportación de una sesión de captura a un libro Excel (.xlsx) organizado.

Hojas generadas:
  * Comentarios — una fila por comentario, un campo por columna.
  * Usuarios    — agregado por autor (nº de comentarios, primera/última vez).
  * Resumen     — metadatos de la sesión y ranking de participantes.
  * Regalos     — solo si la sesión capturó regalos.
"""

from __future__ import annotations

from collections import Counter, OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from capturer import COMMENT_COLUMNS, GIFT_COLUMNS, SESSIONS_DIR

# Paleta (tema claro: fondos claros, texto oscuro)
BRAND = "FE2C55"        # rosa TikTok
BRAND_DARK = "111827"   # casi negro
ACCENT = "25F4EE"       # cian TikTok
SOFT = "FFF1F4"         # rosa muy claro para filas alternas
BORDER = "E5E7EB"

HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
CELL_ALIGN = Alignment(vertical="top", wrap_text=False)
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color=BRAND_DARK)
LABEL_FONT = Font(name="Calibri", size=11, bold=True, color=BRAND_DARK)
VALUE_FONT = Font(name="Calibri", size=11, color="374151")
LINK_FONT = Font(name="Calibri", size=11, color="1155CC", underline="single")
THIN = Side(style="thin", color=BORDER)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Anchos por columna (clave interna -> ancho en caracteres)
WIDTHS = {
    "n": 6, "fecha": 12, "hora": 12, "hora_ms": 16, "timestamp_iso": 26,
    "unix_ms": 17, "latencia_ms": 12, "usuario": 24, "nickname": 24,
    "user_id": 20, "sec_uid": 34, "comentario": 62, "idioma": 9,
    "msg_id": 22, "es_moderador": 11, "es_suscriptor": 11, "es_seguidor": 11,
    "es_mutuo": 11, "es_amigo": 9, "es_top_gifter": 11, "nivel_miembro": 12,
    "nivel_gifter": 12, "insignias": 26, "menciones": 22, "perfil": 34,
    "room_id": 20, "regalo": 22, "gift_id": 12, "diamantes": 12,
    "repeticiones": 12, "combo": 9, "diamantes_totales": 14,
}
CENTERED = {
    "n", "fecha", "hora", "hora_ms", "latencia_ms", "idioma", "es_moderador",
    "es_suscriptor", "es_seguidor", "es_mutuo", "es_amigo", "es_top_gifter",
    "nivel_miembro", "nivel_gifter", "diamantes", "repeticiones", "combo",
    "diamantes_totales", "gift_id",
}


def _write_sheet(ws, columns: list[tuple[str, str]], rows: list[dict], link_key: str | None = None) -> None:
    keys = [k for k, _ in columns]
    ws.append([h for _, h in columns])

    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BOX
    ws.row_dimensions[1].height = 30

    for index, row in enumerate(rows, start=2):
        ws.append([row.get(k, "") for k in keys])
        striped = index % 2 == 0
        for offset, key in enumerate(keys, start=1):
            cell = ws.cell(row=index, column=offset)
            cell.alignment = CELL_ALIGN
            cell.font = VALUE_FONT
            cell.border = BOX
            if striped:
                cell.fill = PatternFill("solid", fgColor=SOFT)
            if key in CENTERED:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            if isinstance(cell.value, int):
                cell.number_format = "#,##0"

    for offset, key in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(offset)].width = WIDTHS.get(key, 18)

    if link_key and link_key in keys:
        col = get_column_letter(keys.index(link_key) + 1)
        for index in range(2, len(rows) + 2):
            cell = ws[f"{col}{index}"]
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.font = LINK_FONT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(keys))}{max(len(rows) + 1, 2)}"


def _summary_sheet(ws, session: Any) -> None:
    status = session.status()
    stats = status["stats"]

    started = session.started_at
    stopped = session.stopped_at
    duration = stats["elapsed_s"]
    hours, remainder = divmod(int(duration), 3600)
    minutes, seconds = divmod(remainder, 60)

    pairs = [
        ("Cuenta de TikTok", f"@{session.unique_id}"),
        ("Sala (room_id)", session.room_id or "—"),
        ("ID de sesión", session.session_id),
        ("Inicio de captura", started.strftime("%Y-%m-%d %H:%M:%S") if started else "—"),
        ("Fin de captura", stopped.strftime("%Y-%m-%d %H:%M:%S") if stopped else "—"),
        ("Duración", f"{hours:02d}:{minutes:02d}:{seconds:02d}"),
        ("Comentarios capturados", stats["comments"]),
        ("Usuarios únicos", stats["unique_users"]),
        ("Comentarios por minuto", stats["per_minute"]),
        ("Regalos capturados", stats["gifts"]),
        ("Exportado", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Fichero CSV", str(session.csv_path)),
    ]

    ws["A1"] = "Resumen de la captura"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")

    for index, (label, value) in enumerate(pairs, start=3):
        ws.cell(row=index, column=1, value=label).font = LABEL_FONT
        cell = ws.cell(row=index, column=2, value=value)
        cell.font = VALUE_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=False)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 52

    # Ranking de participantes por número de comentarios.
    counts = Counter(row.get("usuario", "") for row in session.rows)
    start = len(pairs) + 5
    ws.cell(row=start, column=1, value="Top participantes (por comentarios)").font = TITLE_FONT
    ws.cell(row=start + 1, column=1, value="Usuario").font = LABEL_FONT
    ws.cell(row=start + 1, column=2, value="Comentarios").font = LABEL_FONT
    for offset, (handle, total) in enumerate(counts.most_common(25), start=start + 2):
        ws.cell(row=offset, column=1, value=handle).font = VALUE_FONT
        ws.cell(row=offset, column=2, value=total).font = VALUE_FONT


USER_COLUMNS: list[tuple[str, str]] = [
    ("usuario", "Usuario (@)"),
    ("nickname", "Nombre visible"),
    ("user_id", "ID del autor"),
    ("comentarios", "Comentarios"),
    ("primero", "Primer comentario"),
    ("ultimo", "Último comentario"),
    ("es_suscriptor", "Suscriptor"),
    ("es_moderador", "Moderador"),
    ("es_seguidor", "Seguidor"),
    ("nivel_miembro", "Nivel de miembro"),
    ("perfil", "Perfil"),
]


def export_session(session: Any, output_path: Path | None = None) -> Path:
    """Genera el .xlsx de la sesión y devuelve su ruta."""
    SESSIONS_DIR.mkdir(parents=True, exist_ok=True)
    if output_path is None:
        output_path = SESSIONS_DIR / f"tiktok_{session.unique_id}_{session.session_id}.xlsx"
    output_path = Path(output_path)

    wb = Workbook()

    ws_comments = wb.active
    ws_comments.title = "Comentarios"
    _write_sheet(ws_comments, COMMENT_COLUMNS, session.rows, link_key="perfil")

    _write_sheet(wb.create_sheet("Usuarios"), USER_COLUMNS, sorted(
        _group_users(session.rows),
        key=lambda r: (-r["comentarios"], r["usuario"]),
    ), link_key="perfil")

    if session.gift_rows:
        _write_sheet(wb.create_sheet("Regalos"), GIFT_COLUMNS, session.gift_rows)

    _summary_sheet(wb.create_sheet("Resumen"), session)

    wb.save(output_path)
    return output_path


def _group_users(rows: Iterable[dict]) -> list[dict]:
    grouped: "OrderedDict[str, dict]" = OrderedDict()
    for row in rows:
        handle = row.get("usuario") or "desconocido"
        entry = grouped.get(handle)
        if entry is None:
            grouped[handle] = {
                "usuario": handle,
                "nickname": row.get("nickname", ""),
                "user_id": row.get("user_id", ""),
                "comentarios": 1,
                "primero": row.get("hora_ms", ""),
                "ultimo": row.get("hora_ms", ""),
                "es_suscriptor": row.get("es_suscriptor", ""),
                "es_moderador": row.get("es_moderador", ""),
                "es_seguidor": row.get("es_seguidor", ""),
                "nivel_miembro": row.get("nivel_miembro", ""),
                "perfil": row.get("perfil", ""),
            }
        else:
            entry["comentarios"] += 1
            entry["ultimo"] = row.get("hora_ms", "")
            for flag in ("es_suscriptor", "es_moderador", "es_seguidor"):
                if row.get(flag):
                    entry[flag] = "Sí"
    return list(grouped.values())
