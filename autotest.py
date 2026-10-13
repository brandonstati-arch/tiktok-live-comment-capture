"""
Autoprueba del pipeline completo, sin necesidad de un directo en vivo.

Construye eventos sintéticos con el proto real, los pasa por los handlers de
captura, escribe el CSV, genera el Excel y verifica el contenido resultante.

Uso:  .venv/Scripts/python.exe autotest.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from capturer import COMMENT_COLUMNS, CaptureSession
from demo import feed_demo
from exporter import export_session

FAILURES: list[str] = []


def check(label: str, condition: bool, detail: str = "") -> None:
    mark = "OK  " if condition else "FALLA"
    print(f"  [{mark}] {label}" + (f"  -> {detail}" if detail else ""))
    if not condition:
        FAILURES.append(label)


def main() -> int:
    print("\n== 1. Sesion de captura con eventos sinteticos ==")
    session = CaptureSession("demo_simulado", capture_gifts=True)
    session.state = "live"
    session.room_id = "7300000000000000001"
    session._open_csv()

    feed_demo(session, comments=45, gifts=8)

    check("Se capturaron comentarios", len(session.rows) > 0, f"{len(session.rows)} filas")
    check("Se capturaron regalos", len(session.gift_rows) > 0, f"{len(session.gift_rows)} filas")

    print("\n== 2. Deduplicacion por msg_id ==")
    ids = [r["msg_id"] for r in session.rows]
    check("No hay msg_id duplicados", len(ids) == len(set(ids)),
          f"{len(ids)} filas / {len(set(ids))} ids unicos")

    print("\n== 3. Campos obligatorios de cada comentario ==")
    required = [k for k, _ in COMMENT_COLUMNS]
    sample = session.rows[0]
    missing = [k for k in required if k not in sample]
    check("Todas las columnas definidas existen en la fila", not missing, str(missing))
    check("El comentario no esta vacio", all(r["comentario"].strip() for r in session.rows))
    check("El usuario tiene formato @handle", all(r["usuario"].startswith("@") for r in session.rows))
    check("El ID del autor es numerico", all(isinstance(r["user_id"], int) and r["user_id"] > 0
                                             for r in session.rows))
    check("La hora tiene milisegundos", all(len(r["hora_ms"]) == 12 for r in session.rows),
          session.rows[0]["hora_ms"])
    check("Timestamp ISO presente", all(r["timestamp_iso"] for r in session.rows))
    check("Latencia calculada", all(isinstance(r["latencia_ms"], int) for r in session.rows),
          f"ej. {session.rows[0]['latencia_ms']} ms")
    check("Nº secuencial correcto",
          [r["n"] for r in session.rows] == list(range(1, len(session.rows) + 1)))
    check("Se detecto el moderador",
          any(r["es_moderador"] == "Sí" for r in session.rows))
    check("Se detecto el suscriptor",
          any(r["es_suscriptor"] == "Sí" for r in session.rows))
    check("Se detecto el top gifter",
          any(r["es_top_gifter"] == "Sí" for r in session.rows))
    check("Se detecto el amigo (FRIENDS)",
          any(r["es_amigo"] == "Sí" for r in session.rows))
    check("Insignias legibles extraidas",
          all(r["insignias"] for r in session.rows),
          session.rows[0]["insignias"])
    check("Nivel de gifter (USER_GRADE) presente",
          all(str(r["nivel_gifter"]).isdigit() for r in session.rows),
          f"ej. {session.rows[0]['nivel_gifter']}")
    check("Nivel de miembro (FANS) en suscriptores",
          all(str(r["nivel_miembro"]).isdigit()
              for r in session.rows if r["es_suscriptor"] == "Sí"))

    print("\n== 4. Escritura del CSV ==")
    session._close_files()
    csv_text = session.csv_path.read_text(encoding="utf-8-sig")
    csv_lines = [ln for ln in csv_text.splitlines() if ln.strip()]
    check("CSV creado", session.csv_path.exists(), session.csv_path.name)
    check("CSV con cabecera + filas", len(csv_lines) == len(session.rows) + 1,
          f"{len(csv_lines)} lineas")
    check("Cabecera en espanol", csv_lines[0].startswith("Nº,Fecha,Hora"))

    print("\n== 5. Generacion del Excel ==")
    xlsx = export_session(session)
    check("Excel creado", xlsx.exists(), xlsx.name)

    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    check("Hojas esperadas", set(wb.sheetnames) == {"Comentarios", "Usuarios", "Regalos", "Resumen"},
          str(wb.sheetnames))
    ws = wb["Comentarios"]
    check("Fila de cabecera + datos", ws.max_row == len(session.rows) + 1,
          f"{ws.max_row} filas x {ws.max_column} columnas")
    check("Autofiltro activo", bool(ws.auto_filter.ref), ws.auto_filter.ref)
    check("Panel congelado en A2", ws.freeze_panes == "A2")
    check("Cabecera coloreada", ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("FE2C55"),
          str(ws["A1"].fill.fgColor.rgb))
    check("Columna Comentario con ancho util",
          ws.column_dimensions["L"].width >= 50, str(ws.column_dimensions["L"].width))
    check("Hoja Resumen con metadatos", wb["Resumen"]["A1"].value == "Resumen de la captura")
    check("Hoja Usuarios agregada", wb["Usuarios"].max_row > 1,
          f"{wb['Usuarios'].max_row - 1} usuarios")

    print("\n== 6. Firma de Euler Stream (conexion real, solo lectura) ==")
    try:
        import asyncio

        from TikTokLive import TikTokLiveClient

        async def probe():
            client = TikTokLiveClient(unique_id="@tiktok")
            return await client.is_live()

        result = asyncio.run(probe())
        check("El endpoint de firma/consulta responde", isinstance(result, bool), f"is_live={result}")
    except Exception as exc:  # noqa: BLE001
        check("El endpoint de firma/consulta responde", False, f"{type(exc).__name__}: {exc}")

    print("\n" + "=" * 60)
    if FAILURES:
        print(f"RESULTADO: {len(FAILURES)} comprobacion(es) fallidas -> {FAILURES}")
        return 1
    print("RESULTADO: todas las comprobaciones pasaron.")
    print(f"Excel de ejemplo: {xlsx}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
